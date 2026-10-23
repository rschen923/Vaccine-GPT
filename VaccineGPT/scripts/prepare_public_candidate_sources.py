from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import tarfile
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


DEG_FIELDS = (
    "dataset_id",
    "deg_record_id",
    "gene_symbol",
    "pmid",
    "cog_id",
    "functional_category",
    "product",
    "organism_strain",
    "replicon",
    "condition",
    "reported_phenotype",
    "go_terms",
    "uniprot_accession",
)

AA_ALPHABET = set("ABCDEFGHIKLMNPQRSTVWXYZUO*.")
MISSING_VALUES = {"", "na", "n/a", "null", "[null]", "none", "-", "nan"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_record(
    path: Path, source_url: str, license_status: str, status: str = "verified_local"
) -> dict[str, Any]:
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256_file(path),
        "source_url": source_url,
        "license_status": license_status,
        "status": status,
    }


def audit_normalized_files(output_dir: Path) -> list[dict[str, Any]]:
    records = []
    for path in sorted(output_dir.glob("*.jsonl")):
        with path.open("r", encoding="utf-8") as handle:
            row_count = sum(1 for line in handle if line.strip())
        records.append(
            {
                "path": str(path),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "records": row_count,
            }
        )
    return records


def write_jsonl(rows: Iterable[dict[str, Any]], path: Path) -> int:
    count = 0
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")
            count += 1
    return count


def numeric(value: str | None) -> float | None:
    if value is None or value.strip().casefold() in MISSING_VALUES:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def stable_id(prefix: str, row_number: int, row: dict[str, Any]) -> str:
    payload = json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    digest = hashlib.sha256(f"{row_number}\0{payload}".encode("utf-8")).hexdigest()[:16]
    return f"{prefix}:{row_number}:{digest}"


def is_missing(value: str | None) -> bool:
    return value is None or value.strip().casefold() in MISSING_VALUES


def present_or_none(value: str | None) -> str | None:
    return None if is_missing(value) else value


def prepare_bacfit(path: Path, output: Path) -> dict[str, Any]:
    missing: Counter[str] = Counter()
    taxa: Counter[str] = Counter()
    hosts: Counter[str] = Counter()
    pmids: Counter[str] = Counter()
    raw_row_hashes: Counter[str] = Counter()
    observation_keys: Counter[str] = Counter()
    sequence_hashes: Counter[str] = Counter()
    score_ranges: dict[str, list[float]] = {
        key: [] for key in ("fitness_score", "fitness_z_score", "p_value")
    }
    count = 0

    def records() -> Iterable[dict[str, Any]]:
        nonlocal count
        with zipfile.ZipFile(path) as archive:
            with archive.open("bacfitbase_v1.tsv") as binary:
                text = __import__("io").TextIOWrapper(binary, encoding="utf-8")
                reader = csv.DictReader(text, delimiter="\t")
                for row_number, row in enumerate(reader, 1):
                    count += 1
                    missing.update(
                        key for key, value in row.items() if is_missing(value)
                    )
                    taxa.update([row.get("pathogen_name", "")])
                    hosts.update([row.get("host_name", "")])
                    pmids.update([row.get("pmid", "")])
                    raw_json = json.dumps(row, sort_keys=True, separators=(",", ":"))
                    raw_row_hashes.update([hashlib.sha256(raw_json.encode("utf-8")).hexdigest()])
                    observation = tuple(
                        row.get(key, "")
                        for key in (
                            "pmid",
                            "pathogen_taxon_id",
                            "genbank_locus_tag",
                            "genbank_protein_id",
                            "insertion_site",
                            "fitness_score",
                            "fitness_z_score",
                            "host_taxon_id",
                            "tissue",
                            "time_post_infection",
                        )
                    )
                    observation_keys.update([json.dumps(observation, separators=(",", ":"))])
                    protein = row.get("translation", "").strip().upper()
                    sequence_hashes.update([hashlib.sha256(protein.encode("ascii", "replace")).hexdigest()])
                    for key, values in score_ranges.items():
                        value = numeric(row.get(key))
                        if value is not None:
                            values.append(value)
                    yield {
                        "record_id": stable_id("BacFITBase-v1", row_number, row),
                        "source": "BacFITBase v1",
                        "source_row_number": row_number,
                        "task": "H1",
                        "endpoint": "host-associated infection fitness",
                        "outcome_type": "continuous_measurement",
                        "fitness_score": numeric(row.get("fitness_score")),
                        "fitness_z_score": numeric(row.get("fitness_z_score")),
                        "p_value": numeric(row.get("p_value")),
                        "pathogen_name": row.get("pathogen_name"),
                        "pathogen_taxon_id": present_or_none(row.get("pathogen_taxon_id")),
                        "host_name": row.get("host_name"),
                        "host_taxon_id": present_or_none(row.get("host_taxon_id")),
                        "tissue": row.get("tissue"),
                        "time_post_infection": present_or_none(row.get("time_post_infection")),
                        "locus_tag": row.get("genbank_locus_tag"),
                        "genbank_protein_id": present_or_none(row.get("genbank_protein_id")),
                        "uniprot_accession": present_or_none(row.get("acc")),
                        "gene_symbol": present_or_none(row.get("gene_symbols")),
                        "insertion_site": present_or_none(row.get("insertion_site")),
                        "protein_sequence": protein or None,
                        "label": None,
                        "label_semantics": (
                            "Continuous mutant fitness measurement in the reported infection "
                            "context; not a binary essentiality or vaccine-protection label."
                        ),
                        "source_fields": row,
                    }

    written = write_jsonl(records(), output)
    return {
        "input": str(path),
        "output": str(output),
        "rows": count,
        "rows_written": written,
        "unique_raw_rows": len(raw_row_hashes),
        "exact_duplicate_rows": count - len(raw_row_hashes),
        "unique_context_gene_measurement_keys": len(observation_keys),
        "duplicate_context_gene_measurement_keys": count - len(observation_keys),
        "unique_protein_sequences": len(sequence_hashes),
        "missing_by_field": dict(missing),
        "pathogen_taxa_count": len(taxa),
        "pathogen_taxa": dict(taxa),
        "host_count": len(hosts),
        "hosts": dict(hosts),
        "pmid_count": len(pmids),
        "pmids": dict(pmids),
        "numeric_ranges": {
            key: {"n": len(values), "min": min(values), "max": max(values)}
            for key, values in score_ranges.items()
        },
        "duplicates_retained": True,
    }


def prepare_deg(path: Path, output: Path) -> dict[str, Any]:
    missing: Counter[str] = Counter()
    organisms: Counter[str] = Counter()
    conditions: Counter[str] = Counter()
    raw_row_hashes: Counter[str] = Counter()
    rows = 0

    def records() -> Iterable[dict[str, Any]]:
        nonlocal rows
        with zipfile.ZipFile(path) as archive:
            with archive.open("deg_annotation_p.csv") as binary:
                text = __import__("io").TextIOWrapper(binary, encoding="utf-8-sig")
                reader = csv.reader(text, delimiter=";")
                for row_number, raw in enumerate(reader, 1):
                    if not raw or not any(value.strip() for value in raw):
                        continue
                    values = raw[: len(DEG_FIELDS)]
                    if len(values) != len(DEG_FIELDS):
                        raise ValueError(
                            f"DEG row {row_number} has {len(values)} fields; "
                            f"expected at least {len(DEG_FIELDS)}"
                        )
                    source_fields = dict(zip(DEG_FIELDS, values))
                    rows += 1
                    missing.update(
                        key for key, value in source_fields.items() if is_missing(value)
                    )
                    organisms.update([source_fields["organism_strain"]])
                    conditions.update([source_fields["condition"]])
                    raw_row_hashes.update(
                        [
                            hashlib.sha256(
                                json.dumps(
                                    source_fields, ensure_ascii=False, sort_keys=True
                                ).encode("utf-8")
                            ).hexdigest()
                        ]
                    )
                    yield {
                        "record_id": stable_id("DEG", row_number, source_fields),
                        "source": "DEG annotation",
                        "source_row_number": row_number,
                        "task": "H4",
                        "endpoint": "gene essentiality",
                        "label": 1,
                        "label_semantics": (
                            "Positive record reported in the DEG essential-gene annotation "
                            "table; no unlisted gene is treated as a tested negative."
                        ),
                        "label_scope": "source_dataset_and_reported_condition",
                        "source_fields": source_fields,
                    }

    written = write_jsonl(records(), output)
    return {
        "input": str(path),
        "output": str(output),
        "rows": rows,
        "rows_written": written,
        "unique_raw_rows": len(raw_row_hashes),
        "exact_duplicate_rows": rows - len(raw_row_hashes),
        "organism_strain_count": len(organisms),
        "organism_strains": dict(organisms),
        "condition_count": len(conditions),
        "conditions": dict(conditions),
        "missing_by_field": dict(missing),
        "negative_labels_created": 0,
        "duplicates_retained": True,
    }


def prepare_pa_gold(path: Path, output: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the PLOS supplementary workbook") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    memberships: list[dict[str, Any]] = []
    set_genes: dict[str, set[str]] = {}
    expected = {"GOLD_115": 115, "GOLD_84": 84}
    columns = (
        "locus_tag",
        "old_locus_tag",
        "gene_name",
        "product_description",
        "function_pseudocap",
        "functional_class",
        "localization",
        "accession",
        "go_terms",
    )

    for set_name in expected:
        if set_name not in workbook.sheetnames:
            raise ValueError(f"Missing expected supplementary sheet: {set_name}")
        worksheet = workbook[set_name]
        set_genes[set_name] = set()
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=3, max_col=len(columns), values_only=True), 3
        ):
            if not values or values[0] is None:
                continue
            source_fields = dict(zip(columns, values))
            locus = str(source_fields["locus_tag"]).strip()
            set_genes[set_name].add(locus)
            memberships.append(
                {
                    "record_id": f"PLOS-PCBI-1013945:{set_name}:{locus}",
                    "source": "PLOS Computational Biology 10.1371/journal.pcbi.1013945",
                    "source_row_number": row_number,
                    "task": "H4",
                    "endpoint": "consensus essentiality",
                    "label": 1,
                    "label_semantics": (
                        "Membership in the paper-defined P. aeruginosa essential-gene "
                        "consensus set; GOLD_115 is PA14/LB-specific and GOLD_84 is a "
                        "cross-strain/culture core set."
                    ),
                    "label_scope": set_name,
                    "strain_context": (
                        "PA14 grown in LB" if set_name == "GOLD_115" else
                        "P. aeruginosa consensus across strains and culture conditions"
                    ),
                    "source_fields": source_fields,
                }
            )

    if {key: len(value) for key, value in set_genes.items()} != expected:
        raise ValueError(
            f"Unexpected PLOS set sizes: "
            f"{ {key: len(value) for key, value in set_genes.items()} }"
        )
    written = write_jsonl(memberships, output)
    return {
        "input": str(path),
        "output": str(output),
        "rows_written": written,
        "set_sizes": {key: len(value) for key, value in set_genes.items()},
        "overlap": len(set_genes["GOLD_115"] & set_genes["GOLD_84"]),
        "gold_115_only": len(set_genes["GOLD_115"] - set_genes["GOLD_84"]),
        "gold_84_only": len(set_genes["GOLD_84"] - set_genes["GOLD_115"]),
        "negative_labels_created": 0,
    }


def read_fasta(path: Path) -> Iterable[tuple[int, str, str]]:
    with path.open("r", encoding="utf-8") as handle:
        yield from read_fasta_lines(handle, str(path))


def read_fasta_lines(
    lines: Iterable[str], source: str
) -> Iterable[tuple[int, str, str]]:
    header: str | None = None
    sequence: list[str] = []
    row_number = 0
    for line in lines:
        line = line.strip()
        if line.startswith(">"):
            if header is not None:
                yield row_number, header, "".join(sequence).upper()
            row_number += 1
            header, sequence = line[1:], []
        elif line:
            if header is None:
                raise ValueError(f"Sequence found before the first FASTA header in {source}")
            sequence.append(line)
    if header is not None:
        yield row_number, header, "".join(sequence).upper()


def prepare_t3sepp(path: Path, output_dir: Path) -> dict[str, Any]:
    datasets = (
        (
            "training_positive",
            "T3SEpp_training_testing_datasets/Pos_training_dataset.fasta",
            1,
        ),
        (
            "training_negative",
            "T3SEpp_training_testing_datasets/Neg_training_dataset.fasta",
            0,
        ),
    )
    summaries: dict[str, Any] = {}
    training_sequences: dict[str, int] = {}
    independent_sequences: dict[str, int] = {}
    eib202_records: list[dict[str, Any]] = []
    with tarfile.open(path, mode="r:gz") as archive:
        for split, member_name, label in datasets:
            member = archive.extractfile(member_name)
            if member is None:
                raise ValueError(f"Missing T3SEpp archive member: {member_name}")
            output_path = output_dir / f"t3sepp_{split}.jsonl"
            count = 0
            invalid = 0
            sequence_hashes: Counter[str] = Counter()
            organisms: Counter[str] = Counter()
            with member:
                text = io.TextIOWrapper(member, encoding="utf-8")

                def records() -> Iterable[dict[str, Any]]:
                    nonlocal count, invalid
                    for row_number, header, sequence in read_fasta_lines(text, member_name):
                        count += 1
                        sequence_hash = hashlib.sha256(sequence.encode("ascii", "replace")).hexdigest()
                        sequence_hashes.update([sequence_hash])
                        training_sequences[sequence_hash] = label
                        match = re.search(r"\[([^\]]+)\]\s*$", header)
                        organism = match.group(1) if match else None
                        organisms.update([organism or "unknown"])
                        bad = not sequence or not set(sequence) <= AA_ALPHABET
                        invalid += int(bad)
                        yield {
                            "record_id": f"T3SEpp:{split}:{row_number}",
                            "source": "T3SEpp training/testing benchmark archive",
                            "source_row_number": row_number,
                            "task": "H3_related",
                            "endpoint": "Type III effector sequence classification benchmark",
                            "dataset_split": "training",
                            "benchmark_class": "effector" if label == 1 else "non_effector",
                            "label": label,
                            "label_semantics": (
                                "Positive class in the authors' T3SEpp benchmark; not a "
                                "vaccine-protection or host-protection endpoint."
                                if label == 1
                                else "Negative class in the authors' T3SEpp benchmark; "
                                "benchmark negatives are not evidence that every such "
                                "protein is biologically non-effector."
                            ),
                            "organism": organism,
                            "source_header": header,
                            "protein_sequence": sequence,
                            "sequence_sha256": sequence_hash,
                            "sequence_quality_status": "invalid" if bad else "valid",
                        }

                written = write_jsonl(records(), output_path)
            if count != written:
                raise ValueError(f"T3SEpp {split}: wrote {written} records from {count}")
            summaries[split] = {
                "archive_member": member_name,
                "output": str(output_path),
                "records": count,
                "unique_protein_sequences": len(sequence_hashes),
                "duplicate_protein_sequences": count - len(sequence_hashes),
                "invalid_sequence_records": invalid,
                "organism_value_count_including_unknown": len(organisms),
                "label": label,
            }

        member_name = "T3SEpp_training_testing_datasets/T3SEpp_Independent_Testing_Dataset.xlsx"
        member = archive.extractfile(member_name)
        if member is None:
            raise ValueError(f"Missing T3SEpp archive member: {member_name}")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to read the T3SEpp test workbook") from exc
        workbook = load_workbook(io.BytesIO(member.read()), read_only=True, data_only=True)
        expected_sheets = {"pos-seq-annotation": 42, "neg-seq-annotation": 34}
        output_path = output_dir / "t3sepp_independent_test.jsonl"
        row_count = 0
        sequence_hashes = Counter()
        labels = Counter()
        organisms = Counter()

        def test_records() -> Iterable[dict[str, Any]]:
            nonlocal row_count
            for sheet_name, label in (
                ("pos-seq-annotation", 1),
                ("neg-seq-annotation", 0),
            ):
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(f"Missing T3SEpp test worksheet: {sheet_name}")
                worksheet = workbook[sheet_name]
                header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
                headers = [str(value).strip() for value in header_row]
                expected = {
                    "ID",
                    "gene",
                    "organism",
                    "PMID",
                    "protein accession",
                    "Protein",
                }
                if set(headers) != expected:
                    raise ValueError(f"Unexpected T3SEpp test headers: {headers}")
                for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=2, values_only=True), 2
                ):
                    if not values or not any(value is not None for value in values):
                        continue
                    source_fields = dict(zip(headers, values))
                    sequence = str(source_fields["Protein"] or "").strip().upper()
                    sequence_hash = hashlib.sha256(
                        sequence.encode("ascii", "replace")
                    ).hexdigest()
                    sequence_hashes.update([sequence_hash])
                    independent_sequences[sequence_hash] = label
                    labels.update([label])
                    organisms.update([str(source_fields.get("organism") or "unknown")])
                    row_count += 1
                    if "eib202" in str(source_fields.get("organism") or "").casefold():
                        eib202_records.append(
                            {
                                "gene": source_fields.get("gene"),
                                "label": label,
                                "protein_accession": source_fields.get("protein accession"),
                                "sequence_length": len(sequence),
                            }
                        )
                    yield {
                        "record_id": (
                            f"T3SEpp:independent_test:{sheet_name}:{source_fields['ID']}"
                        ),
                        "source": "T3SEpp independent testing dataset",
                        "source_row_number": row_number,
                        "task": "H3_related",
                        "endpoint": "Type III effector sequence classification benchmark",
                        "dataset_split": "independent_test",
                        "benchmark_class": "effector" if label == 1 else "non_effector",
                        "label": label,
                        "label_semantics": (
                            "Positive class in the authors' independent T3SEpp benchmark; "
                            "not a vaccine-protection endpoint."
                            if label == 1
                            else "Negative class in the authors' independent T3SEpp "
                            "benchmark; does not establish universal non-effector status."
                        ),
                        "organism": source_fields.get("organism"),
                        "pmid": source_fields.get("PMID"),
                        "protein_accession": source_fields.get("protein accession"),
                        "gene": source_fields.get("gene"),
                        "protein_sequence": sequence,
                        "sequence_sha256": sequence_hash,
                        "source_fields": source_fields,
                    }

        written = write_jsonl(test_records(), output_path)
        if row_count != written:
            raise ValueError(f"T3SEpp test: wrote {written} records from {row_count}")
        observed = {
            sheet_name: sum(
                1
                for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
                if row and any(value is not None for value in row)
            )
            for sheet_name in expected_sheets
        }
        if observed != expected_sheets:
            raise ValueError(f"Unexpected T3SEpp test-set sizes: {observed}")
        summaries["independent_test"] = {
            "output": str(output_path),
            "records": row_count,
            "class_counts": dict(labels),
            "unique_protein_sequences": len(sequence_hashes),
            "duplicate_protein_sequences": row_count - len(sequence_hashes),
            "organism_value_count": len(organisms),
            "label": "benchmark labels as supplied; not a protection endpoint",
        }
    overlap = set(training_sequences) & set(independent_sequences)
    summaries["split_integrity"] = {
        "train_test_exact_sequence_overlap": len(overlap),
        "overlap_with_conflicting_class": sum(
            training_sequences[sequence] != independent_sequences[sequence]
            for sequence in overlap
        ),
        "eib202_independent_test_records": eib202_records,
    }
    return {
        "input": str(path),
        "datasets": summaries,
        "license_status": "repository contains a software license; data-specific reuse license not verified",
        "label_review_required_before_m2_h3_training": True,
    }


def prepare_antigendb(path: Path, output: Path) -> dict[str, Any]:
    archive_member = "antigentdb/antigendb_data_full.txt"
    expected_fields = {
        "Aid",
        "Name",
        "Antigen_sequence",
        "DRPDB",
        "Origin",
        "Refer",
        "COM",
        "Synonyms",
    }
    entries: list[dict[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        try:
            raw_text = archive.read(archive_member).decode("utf-8")
        except KeyError as exc:
            raise ValueError(f"Missing AntigenDB archive member: {archive_member}") from exc
    for block in raw_text.split("//"):
        fields: dict[str, str] = {}
        for line in block.splitlines():
            if "\t" not in line:
                if line.strip():
                    raise ValueError(f"Unexpected AntigenDB line without tab: {line[:80]}")
                continue
            key, value = line.split("\t", maxsplit=1)
            fields[key.strip()] = value.strip()
        if not fields:
            continue
        if set(fields) != expected_fields:
            raise ValueError(f"Unexpected AntigenDB fields for Aid={fields.get('Aid')}: {set(fields)}")
        entries.append(fields)

    ids = Counter(entry["Aid"] for entry in entries)
    accessions = Counter(entry["Antigen_sequence"] for entry in entries if entry["Antigen_sequence"])
    comments = Counter(entry["COM"] or "(blank)" for entry in entries)
    source_rows = Counter(
        json.dumps(entry, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        for entry in entries
    )

    def records() -> Iterable[dict[str, Any]]:
        for row_number, fields in enumerate(entries, 1):
            comment = fields["COM"]
            yield {
                "record_id": f"AntigenDB:{fields['Aid']}:{row_number}",
                "source": "AntigenDB curated data archive",
                "source_row_number": row_number,
                "source_database_entry_id": fields["Aid"],
                "task": "H2_related_candidate_evidence",
                "endpoint": "Curated antigen database membership; source study endpoints vary",
                "candidate_status": "curated_antigen_entry",
                "antigenicity_evidence": "reported_in_source_database",
                "protection_status": (
                    "explicitly_not_protective_in_source_comment"
                    if "not protective" in comment.casefold()
                    else "not_specified_by_normalized_source_comment"
                ),
                "organism": fields["Origin"],
                "antigen_name": fields["Name"],
                "protein_accession": present_or_none(fields["Antigen_sequence"]),
                "protein_sequence": None,
                "pdb_accessions": [
                    item.strip()
                    for item in re.split(r"[,;]", fields["DRPDB"])
                    if item.strip()
                ],
                "pmids": [
                    item.strip()
                    for item in re.split(r"[,;]", fields["Refer"])
                    if item.strip()
                ],
                "source_comment": present_or_none(comment),
                "synonyms": present_or_none(fields["Synonyms"]),
                "source_fields": fields,
            }

    written = write_jsonl(records(), output)
    if written != len(entries):
        raise ValueError(f"AntigenDB wrote {written} records from {len(entries)}")
    return {
        "input": str(path),
        "input_archive_member": archive_member,
        "output": str(output),
        "records": written,
        "unique_antigen_ids": len(ids),
        "repeated_source_entry_id_rows": len(entries) - len(ids),
        "exact_duplicate_source_rows": len(entries) - len(source_rows),
        "unique_protein_accessions": len(accessions),
        "records_missing_protein_accession": sum(not entry["Antigen_sequence"] for entry in entries),
        "distinct_source_organism_strings": len({entry["Origin"] for entry in entries}),
        "source_comment_counts": dict(comments),
        "protection_positive_labels_created": 0,
        "protein_sequences_in_source_table": 0,
        "license_status": (
            "Repository MIT license applies to software; README cites database article "
            "CC BY-NC 2.5 UK. Third-party sequence rights remain source-specific."
        ),
    }


def prepare_crispri_tnseq_targets(path: Path, output: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the CRISPRi-TnSeq workbook") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = "Supplementary Table 2"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing CRISPRi-TnSeq worksheet: {sheet_name}")
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    next(rows)
    header = next(rows)
    if tuple(value for value in header[:2]) != (
        "Library of CRISPRi target",
        "IPTG (µM)",
    ):
        raise ValueError(f"Unexpected CRISPRi-TnSeq target headers: {header}")
    targets = []
    for row_number, row in enumerate(rows, 3):
        gene = str(row[0] or "").strip()
        concentrations = str(row[1] or "").strip()
        if not gene:
            continue
        if not concentrations:
            raise ValueError(f"Incomplete CRISPRi-TnSeq target row {row_number}")
        targets.append((gene, concentrations, row_number))
    gene_counts = Counter(gene.casefold() for gene, _, _ in targets)
    if len(targets) != 13 or any(count > 1 for count in gene_counts.values()):
        raise ValueError(f"Unexpected CRISPRi-TnSeq target list: {targets}")
    conditional_gene = "clpP"
    if sum(gene.casefold() == conditional_gene.casefold() for gene, _, _ in targets) != 1:
        raise ValueError("Expected exactly one strain-dependent essential clpP target")

    def records() -> Iterable[dict[str, Any]]:
        for gene, concentrations, row_number in targets:
            conditional = gene.casefold() == conditional_gene.casefold()
            yield {
                "record_id": f"CRISPRi-TnSeq:PMC11371651:{gene}",
                "source": "CRISPRi-TnSeq Supplementary Table 2 and article text",
                "source_row_number": row_number,
                "task": "H4_condition_specific_essentiality_evidence",
                "endpoint": "Paper-reported essentiality status of a CRISPRi target",
                "organism": "Streptococcus pneumoniae D39V",
                "gene_symbol": gene,
                "essentiality_status": (
                    "strain_dependent_essential"
                    if conditional
                    else "reported_essential"
                ),
                "label": None if conditional else 1,
                "label_semantics": (
                    "The paper describes clpP as strain-dependent essential; do not treat "
                    "as universally essential or non-essential."
                    if conditional
                    else "Selected as an essential CRISPRi target in the paper; positive "
                    "evidence only, not a genome-wide negative-labeled screen."
                ),
                "targeting_iptg_concentrations_uM_as_reported": concentrations,
                "experimental_context": (
                    "CRISPRi-TnSeq interaction study in S. pneumoniae D39V; source paper "
                    "states 12 targets were essential and clpP was strain-dependent essential."
                ),
            }

    written = write_jsonl(records(), output)
    if written != len(targets):
        raise ValueError(f"CRISPRi-TnSeq wrote {written} records from {len(targets)}")
    return {
        "input": str(path),
        "input_sheet": sheet_name,
        "output": str(output),
        "records": written,
        "reported_essential_targets": sum(
            gene.casefold() != conditional_gene.casefold() for gene, _, _ in targets
        ),
        "strain_dependent_essential_targets": sum(
            gene.casefold() == conditional_gene.casefold() for gene, _, _ in targets
        ),
        "negative_labels_created": 0,
        "license_status": "Article and supplementary data identify CC BY 4.0; attribution required.",
    }


def prepare_secret6(
    path: Path, output: Path, *, experimentally_validated: bool
) -> dict[str, Any]:
    headers: Counter[str] = Counter()
    sequences: Counter[str] = Counter()
    organisms: Counter[str] = Counter()
    invalid_sequences = 0
    count = 0

    def records() -> Iterable[dict[str, Any]]:
        nonlocal count, invalid_sequences
        for row_number, header, sequence in read_fasta(path):
            count += 1
            identifier = header.split()[0] if header.split() else ""
            headers.update([identifier])
            sequence_hash = hashlib.sha256(sequence.encode("ascii", "replace")).hexdigest()
            sequences.update([sequence_hash])
            match = re.search(r"\[([^\]]+)\]\s*$", header)
            organism = match.group(1) if match else None
            organisms.update([organism or "unknown"])
            invalid = not sequence or not set(sequence) <= AA_ALPHABET
            invalid_sequences += int(invalid)
            yield {
                "record_id": f"SecReT6:{'validated' if experimentally_validated else 'predicted'}:{identifier}:{row_number}",
                "source": "SecReT6 (last update 2021-11-15)",
                "source_row_number": row_number,
                "task": "H3",
                "endpoint": "Type VI secretion system effector",
                "candidate_status": (
                    "experimentally_validated" if experimentally_validated else "bioinformatically_predicted"
                ),
                "label": 1 if experimentally_validated else None,
                "label_semantics": (
                    "Database-designated experimentally validated T6SS effector; "
                    "not a vaccine-protection endpoint."
                    if experimentally_validated
                    else "Predicted sequence candidate; do not use as an experimental positive."
                ),
                "organism": organism,
                "source_header": header,
                "protein_sequence": sequence,
                "sequence_sha256": sequence_hash,
                "sequence_quality_status": "invalid" if invalid else "valid",
            }

    written = write_jsonl(records(), output)
    return {
        "input": str(path),
        "output": str(output),
        "records": count,
        "rows_written": written,
        "unique_identifiers": len(headers),
        "duplicate_identifiers": count - len(headers),
        "unique_protein_sequences": len(sequences),
        "duplicate_protein_sequences": count - len(sequences),
        "invalid_sequence_records": invalid_sequences,
        "organism_count_including_unknown": len(organisms),
        "license_status": "public download; explicit reuse license not located",
        "negative_labels_created": 0,
    }


def prepare_secret6_supplementary(path: Path, output_dir: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the SecReT6 supplementary workbook") from exc

    expected_sheets = {
        "T6SS",
        "T6SE",
        "T6SI",
        "T6SR",
        "T6SA",
        "Recipient susceptibility factor",
    }
    with zipfile.ZipFile(path) as archive:
        workbook = load_workbook(
            io.BytesIO(archive.read("Table_S1.xlsx")), read_only=True, data_only=True
        )
    if set(workbook.sheetnames) != expected_sheets:
        raise ValueError(f"Unexpected SecReT6 Table S1 worksheets: {workbook.sheetnames}")

    summaries: dict[str, Any] = {}
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        headers = [
            str(value).strip() if value is not None else f"unnamed_column_{index}"
            for index, value in enumerate(
                next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True)), 1
            )
        ]
        target = output_dir / (
            "SecReT6_predicted_"
            + re.sub(r"[^A-Za-z0-9]+", "_", sheet_name).strip("_")
            + ".jsonl"
        )
        if target.exists():
            raise FileExistsError(f"Refusing to overwrite existing output: {target}")
        count = 0
        missing: Counter[str] = Counter()
        organisms: Counter[str] = Counter()
        identifiers: Counter[str] = Counter()
        raw_rows: Counter[str] = Counter()

        def records() -> Iterable[dict[str, Any]]:
            nonlocal count
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=4, values_only=True), 4
            ):
                if not values or not any(value is not None for value in values):
                    continue
                source_fields = {
                    header: value
                    for header, value in zip(headers, values)
                    if not header.startswith("unnamed_column_")
                }
                count += 1
                missing.update(
                    key
                    for key in headers
                    if key in source_fields and is_missing(str(source_fields[key]))
                )
                species_value = source_fields.get("species")
                organism_field = "species" if species_value is not None else "strain"
                organism_value = source_fields.get(organism_field)
                if organism_value is not None:
                    organisms.update([str(organism_value)])
                identifier = source_fields.get("protein_id") or source_fields.get("accession")
                if identifier is not None:
                    identifiers.update([str(identifier)])
                raw_rows.update(
                    [
                        hashlib.sha256(
                            json.dumps(
                                source_fields, ensure_ascii=False, sort_keys=True
                            ).encode("utf-8")
                        ).hexdigest()
                    ]
                )
                yield {
                    "record_id": f"SecReT6:TableS1:{sheet_name}:{row_number}",
                    "source": "SecReT6 Table S1 (2021)",
                    "source_row_number": row_number,
                    "candidate_domain": "H3_related",
                    "candidate_category": sheet_name,
                    "candidate_status": "bioinformatically_predicted",
                    "label": None,
                    "label_semantics": (
                        "Database-predicted T6SS-related candidate; not an experimental "
                        "positive or vaccine-protection label."
                    ),
                    "source_fields": source_fields,
                }

        written = write_jsonl(records(), target)
        summaries[sheet_name] = {
            "output": str(target),
            "records": count,
            "rows_written": written,
            "unique_identifiers": len(identifiers),
            "records_with_repeated_identifiers": count - len(identifiers),
            "exact_duplicate_rows": count - len(raw_rows),
            "organism_field": organism_field,
            "organism_value_count": len(organisms),
            "missing_by_field": dict(missing),
            "label_count": 0,
        }
    return {
        "input": str(path),
        "sheets": summaries,
        "license_status": "public download; explicit reuse license not located",
        "rows_labeled": 0,
    }


def audit_ogee_mirror(path: Path) -> dict[str, Any]:
    taxa: Counter[str] = Counter()
    labels: Counter[str] = Counter()
    with path.open("r", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        tax_key = next(
            (
                key
                for key in (reader.fieldnames or [])
                if key.lstrip("#").casefold() in {"taxid", "tax_id"}
            ),
            None,
        )
        label_key = next(
            (
                key
                for key in (reader.fieldnames or [])
                if key.casefold() in {"essential", "essentiality", "label"}
            ),
            None,
        )
        if tax_key is None or label_key is None:
            raise ValueError(f"OGEE mirror schema has no taxID or essentiality field: {reader.fieldnames}")
        for row in reader:
            taxa.update([row.get(tax_key, "")])
            labels.update([row.get(label_key, "")])
        return {
            "input": str(path),
            "columns": reader.fieldnames,
            "taxa": dict(taxa),
            "labels": dict(labels),
            "eligible_bacterial_rows": sum(
                count for taxid, count in taxa.items() if taxid not in {"9606", ""}
            ),
            "decision": "excluded: this mirror contains only human taxID 9606 records",
        }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Normalize public H1/H3/H4 candidate evidence without inventing negatives."
    )
    parser.add_argument(
        "--data-root",
        type=Path,
        default=Path(__file__).resolve().parents[2].parent / "data",
        help="Root data directory containing raw/ and processed/.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="Defaults to processed/public_training_sources_YYYYMMDD_HHMMSS.",
    )
    parser.add_argument(
        "--supplementary-only",
        action="store_true",
        help="Add SecReT6 Table S1 records to an existing generated output directory.",
    )
    parser.add_argument(
        "--t3sepp-only",
        action="store_true",
        help="Add T3SEpp benchmark records to an existing generated output directory.",
    )
    parser.add_argument(
        "--antigendb-only",
        action="store_true",
        help="Add AntigenDB candidate evidence to an existing generated output directory.",
    )
    parser.add_argument(
        "--crispri-tnseq-only",
        action="store_true",
        help="Add reported essential-gene targets from the S. pneumoniae CRISPRi-TnSeq study.",
    )
    args = parser.parse_args()
    raw = args.data_root / "raw" / "public_training_sources" / "reference_databases"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    output_dir = args.output_dir or args.data_root / "processed" / f"public_training_sources_{stamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    selected_incremental_modes = sum(
        (
            args.supplementary_only,
            args.t3sepp_only,
            args.antigendb_only,
            args.crispri_tnseq_only,
        )
    )
    if selected_incremental_modes > 1:
        parser.error("incremental source modes cannot be combined")
    if args.supplementary_only:
        source_audit = output_dir / "source_audit.json"
        if not source_audit.is_file():
            raise FileNotFoundError(f"Expected existing generated audit: {source_audit}")
        supplementary = raw / "SecReT6" / "Table_S1.zip"
        report = json.loads(source_audit.read_text(encoding="utf-8"))
        report["source_files"]["secret6_supplementary"] = file_record(
            supplementary,
            "https://bioinfo-mml.sjtu.edu.cn/SecReT6/download/Table_S1.zip",
            "public download; explicit reuse license not located",
        )
        report["datasets"]["secret6_supplementary"] = prepare_secret6_supplementary(
            supplementary, output_dir
        )
        report["normalized_files"] = audit_normalized_files(output_dir)
        source_audit.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(report["datasets"]["secret6_supplementary"], ensure_ascii=False, indent=2))
        print(f"Audit updated at {source_audit}")
        return
    if args.t3sepp_only:
        source_audit = output_dir / "source_audit.json"
        if not source_audit.is_file():
            raise FileNotFoundError(f"Expected existing generated audit: {source_audit}")
        t3sepp = raw / "T3SEpp" / "T3SEpp_training_testing_datasets.tar.gz"
        report = json.loads(source_audit.read_text(encoding="utf-8"))
        report["source_files"]["t3sepp"] = file_record(
            t3sepp,
            "https://github.com/XinjieHui/T3SEpp/blob/2b7b8814dcf58798fb1f111ede66a6d77b57451e/T3SEpp_training_testing_datasets.tar.gz",
            "repository contains a software license; data-specific reuse license not verified",
        )
        report["datasets"]["t3sepp_h3_benchmark"] = prepare_t3sepp(t3sepp, output_dir)
        report["normalized_files"] = audit_normalized_files(output_dir)
        source_audit.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(report["datasets"]["t3sepp_h3_benchmark"], ensure_ascii=False, indent=2))
        print(f"Audit updated at {source_audit}")
        return
    if args.antigendb_only:
        source_audit = output_dir / "source_audit.json"
        if not source_audit.is_file():
            raise FileNotFoundError(f"Expected existing generated audit: {source_audit}")
        antigendb = raw / "AntigenDB" / "antigendb.zip"
        report = json.loads(source_audit.read_text(encoding="utf-8"))
        report["source_files"]["antigendb"] = file_record(
            antigendb,
            "https://github.com/raghavagps/AntigenDB/blob/f808bac015e0928f6732cb6d67d7cd8892431c93/antigentdb.zip",
            (
                "Repository MIT license applies to software; README cites database article "
                "CC BY-NC 2.5 UK. Third-party sequence rights remain source-specific."
            ),
        )
        report["datasets"]["antigendb"] = prepare_antigendb(
            antigendb, output_dir / "antigendb_curated_antigen_evidence.jsonl"
        )
        report["normalized_files"] = audit_normalized_files(output_dir)
        source_audit.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(report["datasets"]["antigendb"], ensure_ascii=False, indent=2))
        print(f"Audit updated at {source_audit}")
        return
    if args.crispri_tnseq_only:
        source_audit = output_dir / "source_audit.json"
        if not source_audit.is_file():
            raise FileNotFoundError(f"Expected existing generated audit: {source_audit}")
        crispri = raw / "CRISPRi_TnSeq_Spneumoniae_2024" / "41564_2024_1759_MOESM3_ESM.xlsx"
        report = json.loads(source_audit.read_text(encoding="utf-8"))
        report["source_files"]["crispri_tnseq_spneumoniae_2024"] = file_record(
            crispri,
            "https://static-content.springer.com/esm/art%3A10.1038%2Fs41564-024-01759-x/MediaObjects/41564_2024_1759_MOESM3_ESM.xlsx",
            "Article and supplementary data identify CC BY 4.0; attribution required.",
        )
        article_xml = raw / "CRISPRi_TnSeq_Spneumoniae_2024" / "PMC11371651_fullText.xml"
        report["source_files"]["crispri_tnseq_article_full_text"] = file_record(
            article_xml,
            "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC11371651/fullTextXML",
            "CC BY 4.0, as stated in the article full text.",
        )
        report["datasets"]["crispri_tnseq_spneumoniae_2024"] = prepare_crispri_tnseq_targets(
            crispri, output_dir / "crispri_tnseq_spneumoniae_h4_targets.jsonl"
        )
        report["normalized_files"] = audit_normalized_files(output_dir)
        source_audit.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(
            json.dumps(
                report["datasets"]["crispri_tnseq_spneumoniae_2024"],
                ensure_ascii=False,
                indent=2,
            )
        )
        print(f"Audit updated at {source_audit}")
        return
    if any(output_dir.iterdir()):
        raise FileExistsError(f"Refusing to overwrite non-empty output directory: {output_dir}")

    inputs = {
        "bacfitbase_v1": raw / "BacFITBase" / "bacfitbase_v1.zip",
        "deg_annotation": raw / "DEG" / "deg_annotation_p.csv.zip",
        "pa_gold_standard": raw / "PA_gold_standard" / "pcbi.1013945.s011.xlsx",
        "ogee_mirror": raw / "OGEE_mirror_MultiOme" / "OGEE_esential_genes_20190416.txt",
        "secret6_validated": raw / "SecReT6" / "effector_exp_protein.fasta",
        "secret6_predicted": raw / "SecReT6" / "effector_pre_protein.fasta",
        "secret6_supplementary": raw / "SecReT6" / "Table_S1.zip",
        "t3sepp": raw / "T3SEpp" / "T3SEpp_training_testing_datasets.tar.gz",
    }
    missing = [str(path) for path in inputs.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing expected public-source files:\n" + "\n".join(missing))

    report: dict[str, Any] = {
        "schema_version": "1.0",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "output_dir": str(output_dir),
        "source_files": {
            "bacfitbase_v1": file_record(
                inputs["bacfitbase_v1"],
                "https://bacfitbase.tartaglialab.com/downloads/bacfitbase_v1.zip",
                "public source; explicit data reuse license not verified",
            ),
            "deg_annotation": file_record(
                inputs["deg_annotation"],
                "http://tubic.org/deg/public/download/deg_annotation_p.csv.zip",
                "public database; explicit data reuse license not verified",
            ),
            "pa_gold_standard": file_record(
                inputs["pa_gold_standard"],
                "https://doi.org/10.1371/journal.pcbi.1013945",
                "PLOS article and supplementary data are CC BY 4.0",
            ),
            "ogee_mirror": file_record(
                inputs["ogee_mirror"],
                "https://github.com/menchelab/MultiOme",
                "mirror license and provenance not verified",
            ),
            "secret6_validated": file_record(
                inputs["secret6_validated"],
                "https://bioinfo-mml.sjtu.edu.cn/SecReT6/download/effector_exp_protein.fasta",
                "public download; explicit reuse license not located",
            ),
            "secret6_predicted": file_record(
                inputs["secret6_predicted"],
                "https://bioinfo-mml.sjtu.edu.cn/SecReT6/download/effector_pre_protein.fasta",
                "public download; explicit reuse license not located",
            ),
            "secret6_supplementary": file_record(
                inputs["secret6_supplementary"],
                "https://bioinfo-mml.sjtu.edu.cn/SecReT6/download/Table_S1.zip",
                "public download; explicit reuse license not located",
            ),
            "t3sepp": file_record(
                inputs["t3sepp"],
                "https://github.com/XinjieHui/T3SEpp/blob/2b7b8814dcf58798fb1f111ede66a6d77b57451e/T3SEpp_training_testing_datasets.tar.gz",
                "repository contains a software license; data-specific reuse license not verified",
            ),
        },
        "datasets": {},
        "processing_rules": [
            "Raw source files are read-only and are never rewritten.",
            "Repeated measurements and source rows are retained; duplicate statistics are reported.",
            "No unlisted, unmeasured, or unmatched gene is converted to a negative label.",
            "BacFITBase fitness values remain continuous and retain infection context.",
            "SecReT6 predicted candidates are separate from experimentally validated records.",
            "OGEE mirror records are excluded from bacterial datasets when taxID is human.",
        ],
    }

    report["datasets"]["bacfitbase_h1"] = prepare_bacfit(
        inputs["bacfitbase_v1"], output_dir / "bacfitbase_h1_observations.jsonl"
    )
    report["datasets"]["deg_h4"] = prepare_deg(
        inputs["deg_annotation"], output_dir / "deg_h4_positive_observations.jsonl"
    )
    report["datasets"]["pa_gold_h4"] = prepare_pa_gold(
        inputs["pa_gold_standard"], output_dir / "pa_gold_h4_memberships.jsonl"
    )
    report[    "datasets"]["secreT6_h3_validated"] = prepare_secret6(
        inputs["secret6_validated"],
        output_dir / "SecReT6_h3_validated_t6se.jsonl",
        experimentally_validated=True,
    )
    report[    "datasets"]["secreT6_h3_predicted"] = prepare_secret6(
        inputs["secret6_predicted"],
        output_dir / "SecReT6_h3_predicted_candidates.jsonl",
        experimentally_validated=False,
    )
    report["datasets"]["secret6_supplementary"] = prepare_secret6_supplementary(
        inputs["secret6_supplementary"], output_dir
    )
    report["datasets"]["t3sepp_h3_benchmark"] = prepare_t3sepp(
        inputs["t3sepp"], output_dir
    )
    report["datasets"]["ogee_mirror"] = audit_ogee_mirror(inputs["ogee_mirror"])
    report["normalized_files"] = audit_normalized_files(output_dir)
    (output_dir / "source_audit.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(report["datasets"], ensure_ascii=False, indent=2))
    print(f"Audit written to {output_dir / 'source_audit.json'}")


if __name__ == "__main__":
    main()
